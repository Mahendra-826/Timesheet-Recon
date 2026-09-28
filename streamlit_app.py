import json
import os
import subprocess
import sys
import calendar
import re
from pathlib import Path
import streamlit as st
import pandas as pd
from openpyxl import load_workbook
import streamlit.components.v1 as components
from config import (
    INPUT_DIR,
    CLIENT_TIMESHEET,
    PARENT_TIMESHEET,
    OUTPUT_FILE,
    PENDING_REVIEW_FILE,
    GEMINI_MODEL,
)
from hours_governance import apply_pending_corrections





def get_parent_month(parent_path):
    """Return a three-letter month from the parent filename or workbook."""

    month_names = {
        name.lower(): name[:3]
        for name in calendar.month_name
        if name
    }
    month_names.update({
        name.lower(): name[:3]
        for name in calendar.month_abbr
        if name
    })

    filename = Path(parent_path).stem.lower()
    for month, abbreviation in month_names.items():
        if re.search(rf"(?<![a-z]){re.escape(month)}(?![a-z])", filename):
            return abbreviation

    try:
        preview = pd.read_excel(parent_path, header=None, nrows=10)
        for value in preview.astype(str).stack():
            match = re.search(
                r"\b(" + "|".join(month_names) + r")\b",
                value,
                flags=re.IGNORECASE,
            )
            if match:
                return month_names[match.group(1).lower()]
    except Exception:
        pass

    return "Unknown"


st.title("Timesheet Governance UI")
st.markdown("Upload the timesheets below, then run governance and download the output.")

# --------------------------------------------------------------------
# Session state
#
# Streamlit reruns this whole script on every widget interaction, so
# anything that must survive past the "Run Governance" button click
# (whether processing finished, which weeks still need a human-
# confirmed Billable Hours value, whether that review is done) has to
# live in st.session_state rather than a local variable.
# --------------------------------------------------------------------

if "processing_done" not in st.session_state:
    st.session_state.processing_done = False

if "pending_items" not in st.session_state:
    st.session_state.pending_items = []

if "reviews_resolved" not in st.session_state:
    st.session_state.reviews_resolved = False

if "run_parent_path" not in st.session_state:
    st.session_state.run_parent_path = None

# --------------------------------------------------------------------
# Gemini configuration (sidebar)
#
# For deployment, the API key/model shouldn't have to live in the
# server's venv/.env — the user provides them here instead. Left
# blank, the run falls back to whatever is configured server-side
# (config.py / .env), so this stays optional for local/dev use.
# --------------------------------------------------------------------

st.sidebar.header("Gemini Configuration")

gemini_api_key_input = st.sidebar.text_input(
    "Gemini API Key",
    type="password",
    help="Used only for this run. Leave blank to use the server's configured key.",
)

gemini_model_input = st.sidebar.text_input(
    "Gemini Model",
    value=GEMINI_MODEL,
    help="e.g. gemini-2.5-flash",
)

client_file = st.file_uploader("Upload Client Timesheet (Expedia) (.xlsx)", type=["xlsx"], key="client")
parent_file = st.file_uploader("Upload Parent Timesheet (Tenarai) (.xlsx)", type=["xlsx"], key="parent")

# Log mode dropdown: Final only (default) or Runtime streaming
log_mode = st.selectbox("Log mode", ["Final only", "Runtime (stream)"], index=0)

# Save uploads immediately using their original filenames so main.py can use them
saved_client_path = None
saved_parent_path = None

if client_file is not None:
    try:
        INPUT_DIR.mkdir(parents=True, exist_ok=True)
        saved_client_path = Path(INPUT_DIR) / client_file.name
        with open(saved_client_path, "wb") as f:
            f.write(client_file.getbuffer())
        # saved silently
    except Exception as e:
        st.warning(f"Could not save client upload immediately: {e}")

if parent_file is not None:
    try:
        INPUT_DIR.mkdir(parents=True, exist_ok=True)
        saved_parent_path = Path(INPUT_DIR) / parent_file.name
        with open(saved_parent_path, "wb") as f:
            f.write(parent_file.getbuffer())
        # saved silently
    except Exception as e:
        st.warning(f"Could not save parent upload immediately: {e}")
detected_info = None

if st.button("Run Governance"):
    # determine paths to run with
    if saved_client_path is None and client_file is None:
        st.error("Please upload a client timesheet.")
        st.stop()

    if saved_parent_path is None and parent_file is None:
        st.error("Please upload a parent timesheet.")
        st.stop()

    # Use saved paths when available, otherwise fall back to configured paths
    run_client = saved_client_path if saved_client_path is not None else Path(CLIENT_TIMESHEET)
    run_parent = saved_parent_path if saved_parent_path is not None else Path(PARENT_TIMESHEET)

    with st.spinner("Running main.py..."):
        # Run main.py in unbuffered mode so live prints appear promptly
        cmd = [sys.executable, "-u", "main.py", str(run_client), str(run_parent)]

        # Override the subprocess's Gemini API key/model with whatever
        # was entered in the sidebar for this run; fall back to the
        # server's own environment when a field is left blank.
        run_env = os.environ.copy()

        if gemini_api_key_input:
            run_env["GEMINI_API_KEY"] = gemini_api_key_input

        if gemini_model_input:
            run_env["GEMINI_MODEL"] = gemini_model_input

        proc = subprocess.Popen(
            cmd,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            bufsize=1,
            cwd=Path(__file__).parent,
            env=run_env,
        )

        logs = []

        if log_mode == "Runtime (stream)":
            # Stream logs live into a code box while the process runs
            log_box = st.empty()

            try:
                while True:
                    line = proc.stdout.readline()
                    if line:
                        logs.append(line)
                        # show recent logs (limit size to avoid huge payload)
                        log_box.code("".join(logs[-1000:]))
                    elif proc.poll() is not None:
                        break

                returncode = proc.returncode

                # ensure any remaining output is captured
                remaining = proc.stdout.read()
                if remaining:
                    logs.append(remaining)

            except Exception as e:
                logs.append(f"\n[STREAMING ERROR] {e}\n")
                returncode = proc.poll() or 1

        else:
            # Final-only mode: wait and capture all output at the end
            out, _ = proc.communicate()
            returncode = proc.returncode
            logs = out.splitlines(keepends=True) if out is not None else []

        # Show logs and handle errors
        #
        # In "Runtime (stream)" mode, every line was already streamed
        # live above -- showing the same full log again afterward
        # would just be a redundant duplicate. Only "Final only" mode
        # (which streams nothing while the process runs) needs it as
        # the sole way to see what happened.
        full_logs = "".join(logs)
        if returncode != 0:
            st.error("Processing failed. See logs below.")

            if log_mode != "Runtime (stream)":
                with st.expander("Process logs (full)", expanded=False):
                    st.code(full_logs)

            if "503" in full_logs or "UNAVAILABLE" in full_logs:
                st.warning("The model/service returned a 503 (high demand). Try again later or rerun.")

            st.session_state.processing_done = False
        else:
            st.success("Processing completed.")

            if log_mode != "Runtime (stream)":
                with st.expander("Process logs (full)", expanded=False):
                    st.code(full_logs)

            # ----------------------------------------------------------------
            # Load any weeks HoursGovernance flagged for manual review
            # (duplicate employee/week rows with an offsetting +/-
            # Billable Hours pair) before we show the final report.
            # ----------------------------------------------------------------

            if PENDING_REVIEW_FILE.exists():
                with open(PENDING_REVIEW_FILE, "r", encoding="utf-8") as f:
                    pending_items = json.load(f)
            else:
                pending_items = []

            st.session_state.processing_done = True
            st.session_state.pending_items = pending_items
            st.session_state.reviews_resolved = len(pending_items) == 0
            st.session_state.run_parent_path = str(run_parent)


# --------------------------------------------------------------------
# Review step
#
# Shown only when governance flagged one or more weeks with a
# duplicate-row correction/reversal pattern (e.g. +40 then -40). The
# pipeline already finished and produced a report, but those specific
# weeks were colored "pending review" instead of Green/Red until a
# human confirms the real Billable Hours. Every flagged week is
# reviewed together in one form, not one at a time.
# --------------------------------------------------------------------

if st.session_state.processing_done and not st.session_state.reviews_resolved:

    pending_items = st.session_state.pending_items

    st.warning(
        f"{len(pending_items)} week(s) have a duplicate-entry "
        "correction (e.g. hours entered, then reversed) and need you "
        "to confirm the final Billable Hours before the report is "
        "finalized."
    )

    with st.form("pending_review_form"):

        confirmed_values = {}

        for index, item in enumerate(pending_items):

            st.markdown(
                f"**{item['parent_name']}** "
                f"({item['week_start']} to {item['week_end']})"
            )

            st.caption(
                f"Timesheet Status: {item['client_status'] or '(blank)'} "
                f"— duplicate Billable Hours entries found: "
                f"{item['original_values']}"
            )

            confirmed_values[item["key"]] = st.number_input(
                "Confirmed Billable Hours",
                value=float(item["suggested_value"]),
                key=f"pending_review_{index}",
            )

        submitted = st.form_submit_button("Apply corrections & continue")

    if submitted:
        apply_pending_corrections(confirmed_values)
        st.session_state.reviews_resolved = True
        st.rerun()

# --------------------------------------------------------------------
# Final preview + download
#
# Only shown once processing has finished AND every flagged week (if
# any) has been resolved by the review step above.
# --------------------------------------------------------------------

if st.session_state.processing_done and st.session_state.reviews_resolved:

    if OUTPUT_FILE.exists():
        # Provide a full-sheet styled HTML preview that preserves cell background colors
        try:
            def generate_colored_html(path, max_rows=None, max_cols=None, max_height=800):
                wb = load_workbook(path, data_only=True)
                ws = wb.active
                max_row = ws.max_row if max_rows is None else min(ws.max_row, max_rows)
                max_col = ws.max_column if max_cols is None else min(ws.max_column, max_cols)

                html_parts = [
                    f'<div style="overflow:auto; max-width:100%; max-height:{max_height}px;">',
                    '<table border="1" cellspacing="0" cellpadding="4" style="border-collapse:collapse; background-color:#FFFFFF;">',
                ]

                for r in range(1, max_row + 1):
                    html_parts.append('<tr>')
                    for c in range(1, max_col + 1):
                        cell = ws.cell(row=r, column=c)
                        val = cell.value if cell.value is not None else ''
                        bg = None
                        try:
                            col_obj = cell.fill.start_color
                            if col_obj is not None:
                                rgb = getattr(col_obj, 'rgb', None)
                                if rgb and len(rgb) in (6, 8):
                                    if len(rgb) == 8:
                                        rgb = rgb[-6:]
                                    # ignore default/empty black fills
                                    if rgb.lower() not in ("000000", "00000000"):
                                        bg = f'#{rgb}'
                        except Exception:
                            bg = None

                        # default white background and black text for readability
                        style = 'background-color:#FFFFFF;color:#000000;'
                        if bg:
                            style = f'background-color:{bg};color:#000000;'

                        safe_val = str(val)
                        html_parts.append(f'<td style="{style}">{safe_val}</td>')
                    html_parts.append('</tr>')

                html_parts.append('</table></div>')
                return ''.join(html_parts)

            # Render with a reasonable height and allow internal scrolling
            max_preview_height = 400
            html = generate_colored_html(OUTPUT_FILE, max_height=max_preview_height)
            components.html(html, height=max_preview_height + 50)
        except Exception as e:
            st.info(f"Could not render colored preview: {e}")
        with open(OUTPUT_FILE, "rb") as f:
            data = f.read()

        st.download_button(
            label="Download Output Excel",
            data=data,
            file_name=f"Reco_{get_parent_month(st.session_state.run_parent_path)}_file.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    else:
        st.error(f"Expected output file not found: {OUTPUT_FILE}")

